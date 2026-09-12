import os
import csv
import fitz  # PyMuPDF
from docx import Document
import openpyxl


def extract_from_pdf(file_path):
    """Extract text from a PDF file."""
    text = ""

    try:
        document = fitz.open(file_path)

        for page in document:
            text += page.get_text()

        document.close()
        return text.strip()

    except Exception as e:
        raise RuntimeError(f"Unable to extract PDF: {e}")


def extract_from_docx(file_path):
    """Extract text from a DOCX file."""
    try:
        document = Document(file_path)

        text = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)

        return "\n".join(text).strip()

    except Exception as e:
        raise RuntimeError(f"Unable to extract DOCX: {e}")


def extract_from_xlsx(file_path):
    """Extract data from an Excel file."""

    try:
        workbook = openpyxl.load_workbook(
            file_path,
            data_only=True
        )

        all_data = []

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    all_data.append(" | ".join(values))

        workbook.close()

        return "\n".join(all_data).strip()

    except Exception as e:
        raise RuntimeError(f"Unable to extract XLSX: {e}")


def extract_from_csv(file_path):
    """Extract data from a CSV file."""

    try:
        rows = []

        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.reader(file)

            for row in reader:
                values = [
                    str(value).strip()
                    for value in row
                    if value.strip()
                ]

                if values:
                    rows.append(" | ".join(values))

        return "\n".join(rows).strip()

    except Exception as e:
        raise RuntimeError(f"Unable to extract CSV: {e}")


def extract_content(file_path):
    """
    Automatically select the correct extraction method
    based on the file extension.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    elif extension == ".docx":
        return extract_from_docx(file_path)

    elif extension == ".xlsx":
        return extract_from_xlsx(file_path)

    elif extension == ".csv":
        return extract_from_csv(file_path)

    else:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )


if __name__ == "__main__":

    print("Document Extractor Module")
    print("=" * 40)

    print("Supported formats:")
    print("1. PDF")
    print("2. DOCX")
    print("3. XLSX")
    print("4. CSV")

    print("\nExtractor module loaded successfully.")